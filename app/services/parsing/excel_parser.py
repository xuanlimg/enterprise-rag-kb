from pathlib import Path

from openpyxl import load_workbook

from app.services.parsing.base import BaseParser, ParsedDocument, ParsedPage


class ExcelParser(BaseParser):
    """Excel 解析器，按工作表展开单元格文本。"""

    def parse(self, path: Path) -> ParsedDocument:
        workbook = load_workbook(str(path), data_only=True, read_only=True)
        pages: list[ParsedPage] = []

        for sheet in workbook.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]
                if values:
                    lines.append(" | ".join(values))
            if lines:
                pages.append(
                    ParsedPage(
                        page_no=None,
                        text=f"# Sheet: {sheet.title}\n" + "\n".join(lines),
                    )
                )

        full_text = "\n\n".join(page.text for page in pages)
        return ParsedDocument(
            text=full_text,
            pages=pages,
            metadata={"parser": "excel", "sheet_count": len(workbook.worksheets)},
        )
